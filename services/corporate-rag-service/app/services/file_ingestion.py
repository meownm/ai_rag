from __future__ import annotations

import hashlib
import io
from pathlib import Path

from app.services.ingestion import SourceItem


class FileByteIngestor:
    SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}

    def ingest_bytes(self, *, filename: str, payload: bytes) -> SourceItem:
        ext = Path(filename).suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file extension: {ext}")

        markdown = self._to_markdown(ext, payload)
        digest = hashlib.sha256(payload).hexdigest()
        return SourceItem(
            source_type="FILE_UPLOAD_OBJECT",
            external_ref=digest,
            title=filename,
            markdown=markdown,
            url="",
            labels=["upload", ext.lstrip(".")],
        )

    def _to_markdown(self, ext: str, payload: bytes) -> str:
        if ext in {".txt", ".md"}:
            return payload.decode("utf-8", errors="ignore")
        if ext == ".docx":
            return self._docx_to_markdown(payload)
        if ext == ".pdf":
            return self._pdf_to_markdown(payload)
        raise ValueError(f"Unsupported file extension: {ext}")

    def _docx_to_markdown(self, payload: bytes) -> str:
        import docx

        document = docx.Document(io.BytesIO(payload))
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                lines.append(f"{'#' * max(1, int(level))} {text}")
            else:
                lines.append(text)

        for table in document.tables:
            if not table.rows:
                continue
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header = rows[0]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
        return "\n\n".join(lines)

    def _pdf_to_markdown(self, payload: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(payload))
        pages: list[str] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"<!-- page:{idx} -->\n{text}")
        return "\n\n".join(pages)
